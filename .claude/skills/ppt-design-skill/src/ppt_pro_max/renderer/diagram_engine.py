"""DiagramEngine — main entry point for structured diagram rendering."""

from __future__ import annotations

from typing import Any

from ppt_pro_max.renderer.diagram.base import BaseDiagram
from ppt_pro_max.renderer.diagram.diagram_style import DiagramStyle
from ppt_pro_max.renderer.diagram.layout_engine import Region


class DiagramEngine:

    _registry: dict[str, type[BaseDiagram]] = {}

    @classmethod
    def register(cls, diagram_type: str, diagram_cls: type[BaseDiagram]) -> None:
        cls._registry[diagram_type] = diagram_cls

    @classmethod
    def get_supported_types(cls) -> list[str]:
        return sorted(cls._registry.keys())

    def render(
        self,
        slide,
        diagram_type: str,
        data: dict[str, Any],
        style: DiagramStyle,
        region: Region,
    ) -> None:
        diagram_cls = self._registry.get(diagram_type)
        if diagram_cls is None:
            self._render_as_text_list(slide, diagram_type, data, style, region)
            return
        diagram = diagram_cls(data, style, region)
        diagram.render(slide)

    def _render_as_text_list(
        self,
        slide,
        diagram_type: str,
        data: dict[str, Any],
        style: DiagramStyle,
        region: Region,
    ) -> None:
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor

        items = []
        nodes = data.get("nodes", [])
        if nodes:
            for node in nodes:
                label = node.get("label", node.get("title", ""))
                if label:
                    items.append(label)
        stages = data.get("stages", [])
        if stages:
            for stage in stages:
                label = stage.get("label", stage.get("name", ""))
                if label:
                    items.append(label)
        events = data.get("events", [])
        if events:
            for event in events:
                label = event.get("label", event.get("date", ""))
                if label:
                    items.append(label)

        if not items:
            items = [diagram_type.replace("_", " ").title()]

        left = Inches(region.left + 0.2)
        top = Inches(region.top + 0.2)
        width = Inches(max(region.width - 0.4, 1.0))
        height = Inches(max(region.height - 0.4, 1.0))

        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True

        for i, item in enumerate(items):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = f"• {item}"
            font_color = style.resolve_color(style.node_font_color)
            p.font.size = Pt(style.node_font_size_pt)
            p.font.color.rgb = RGBColor.from_string(font_color.lstrip("#"))


def _auto_register():
    try:
        from ppt_pro_max.renderer.diagram.flowchart import FlowchartDiagram
        DiagramEngine.register("flowchart", FlowchartDiagram)
    except ImportError:
        pass
    try:
        from ppt_pro_max.renderer.diagram.funnel import FunnelDiagram
        DiagramEngine.register("funnel", FunnelDiagram)
    except ImportError:
        pass
    try:
        from ppt_pro_max.renderer.diagram.timeline import TimelineDiagram
        DiagramEngine.register("timeline", TimelineDiagram)
    except ImportError:
        pass
    try:
        from ppt_pro_max.renderer.diagram.swot import SwotDiagram
        DiagramEngine.register("swot", SwotDiagram)
    except ImportError:
        pass
    try:
        from ppt_pro_max.renderer.diagram.matrix import MatrixDiagram
        DiagramEngine.register("matrix", MatrixDiagram)
    except ImportError:
        pass
    try:
        from ppt_pro_max.renderer.diagram.cycle import CycleDiagram
        DiagramEngine.register("cycle", CycleDiagram)
    except ImportError:
        pass
    try:
        from ppt_pro_max.renderer.diagram.table import TableDiagram
        DiagramEngine.register("table", TableDiagram)
    except ImportError:
        pass
    try:
        from ppt_pro_max.renderer.diagram.hierarchy import HierarchyDiagram
        DiagramEngine.register("hierarchy", HierarchyDiagram)
    except ImportError:
        pass
    try:
        from ppt_pro_max.renderer.diagram.pyramid import PyramidDiagram
        DiagramEngine.register("pyramid", PyramidDiagram)
    except ImportError:
        pass
    try:
        from ppt_pro_max.renderer.diagram.venn import VennDiagram
        DiagramEngine.register("venn", VennDiagram)
    except ImportError:
        pass


_auto_register()
